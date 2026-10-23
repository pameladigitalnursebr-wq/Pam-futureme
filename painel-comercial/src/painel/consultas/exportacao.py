"""Exportação dos resultados para CSV e Excel, com os mesmos filtros do painel."""

from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any

import psycopg
from openpyxl import Workbook

from . import metricas
from .filtros import DIMENSOES, Filtros

LIMITE_LINHAS = 200_000
_POR_PAGINA = 1000


def _linhas_detalhe(conn: psycopg.Connection, f: Filtros, dados_pessoais: bool) -> tuple[list[str], list[list[Any]]]:
    primeira = metricas.detalhe(conn, f, 1, _POR_PAGINA, "recentes", dados_pessoais)
    colunas = list(primeira["colunas"].keys())
    linhas = [[i.get(c) for c in colunas] for i in primeira["itens"]]
    paginas = min(primeira["paginas"], LIMITE_LINHAS // _POR_PAGINA)
    for p in range(2, paginas + 1):
        linhas += [[i.get(c) for c in colunas] for i in metricas.detalhe(conn, f, p, _POR_PAGINA, "recentes",
                                                                           dados_pessoais)["itens"]]
    return list(primeira["colunas"].values()), linhas


def _linhas_ranking(conn: psycopg.Connection, dimensao: str, f: Filtros) -> tuple[list[str], list[list[Any]]]:
    r = metricas.ranking(conn, dimensao, f, limite=500)
    cabecalho = [DIMENSOES[dimensao].titulo, "Alunos", "Interesses", "% dos alunos do recorte"]
    return cabecalho, [[i["rotulo"], i["alunos"], i["interesses"], i["percentual"]] for i in r["itens"]]


def _descricao_filtros(conn: psycopg.Connection, f: Filtros) -> list[list[str]]:
    linhas = []
    for nome, vals in f.valores.items():
        dim = DIMENSOES[nome]
        nomes = metricas.rotulos(conn, dim, vals)
        linhas.append([dim.titulo, "; ".join(str(nomes.get(v, v)) for v in vals)])
    return linhas or [["Filtros", "nenhum (base inteira)"]]


def exportar(conn: psycopg.Connection, f: Filtros, formato: str, tipo: str = "detalhe",
             dimensao: str | None = None, dados_pessoais: bool = False) -> tuple[bytes, str, str]:
    """Devolve (conteúdo, media type, nome do arquivo)."""
    if tipo == "ranking":
        if dimensao not in DIMENSOES:
            raise ValueError("Informe a dimensão do ranking.")
        cabecalho, linhas = _linhas_ranking(conn, dimensao, f)
        base = f"ranking_{dimensao}"
    else:
        cabecalho, linhas = _linhas_detalhe(conn, f, dados_pessoais)
        base = "alunos"
    carimbo = datetime.now().strftime("%Y%m%d_%H%M")
    if formato == "xlsx":
        wb = Workbook(write_only=True)
        dados = wb.create_sheet("Dados")
        dados.append(cabecalho)
        for linha in linhas:
            dados.append([v.replace(tzinfo=None) if isinstance(v, datetime) else v for v in linha])
        info = wb.create_sheet("Filtros")
        info.append(["Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")])
        for linha in _descricao_filtros(conn, f):
            info.append(linha)
        buf = io.BytesIO()
        wb.save(buf)
        return (buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                f"{base}_{carimbo}.xlsx")
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")    # ; abre direto no Excel em português
    w.writerow(cabecalho)
    w.writerows(linhas)
    return buf.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", f"{base}_{carimbo}.csv"
